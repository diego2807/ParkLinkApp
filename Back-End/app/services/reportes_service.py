import io
from datetime import datetime

# Importaciones para la generación de Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Importaciones para la generación de PDF profesional
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


class ReportesService:

    @staticmethod
    def generar_excel_completo(datos):
        """
        Genera un archivo Excel (.xlsx) estructurado en pestañas para todo el historial
        del aplicativo: Resumen, Usuarios, Vehículos, Historial de Accesos y Reservas.
        """
        wb = Workbook()

        # Estilos visuales estándar
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid") # Azul ParkLink
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1A365D")
        border_thin = Border(
            left=Side(style='thin', color='CBD5E0'),
            right=Side(style='thin', color='CBD5E0'),
            top=Side(style='thin', color='CBD5E0'),
            bottom=Side(style='thin', color='CBD5E0')
        )

        # ---------------------------------------------------------
        # 1. PESTAÑA: RESUMEN GENERAL
        # ---------------------------------------------------------
        ws_resumen = wb.active
        ws_resumen.title = "Resumen General"
        
        ws_resumen.append(["PARKLINK - REPORTE CONSOLIDADO DEL SISTEMA"])
        ws_resumen.append([f"Fecha de generación: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        ws_resumen.append([]) # Fila en blanco
        
        ws_resumen.append(["Métrica del Sistema", "Cantidad Total"])
        ws_resumen.append(["Usuarios Registrados", len(datos.get('usuarios', []))])
        ws_resumen.append(["Vehículos Registrados", len(datos.get('vehiculos', []))])
        ws_resumen.append(["Historial de Accesos / Movimientos", len(datos.get('accesos', []))])
        ws_resumen.append(["Reservas Procesadas", len(datos.get('reservas', []))])

        ws_resumen['A1'].font = title_font

        # ---------------------------------------------------------
        # 2. PESTAÑA: HISTORIAL DE ACCESOS
        # ---------------------------------------------------------
        ws_accesos = wb.create_sheet(title="Historial Accesos")
        ws_accesos.append(["ID", "Usuario", "Documento", "Placa", "Tipo Vehículo", "Fecha Entrada", "Fecha Salida", "Estado"])
        
        for acc in datos.get('accesos', []):
            ws_accesos.append([
                acc.get('id', 'N/A'),
                acc.get('usuario', 'N/A'),
                acc.get('documento', 'N/A'),
                acc.get('placa', 'N/A'),
                acc.get('tipo_vehiculo', 'N/A'),
                str(acc.get('fecha_entrada', 'N/A')),
                str(acc.get('fecha_salida', 'En Parqueadero')) if acc.get('fecha_salida') else 'En Parqueadero',
                acc.get('estado', 'N/A')
            ])

        # ---------------------------------------------------------
        # 3. PESTAÑA: USUARIOS
        # ---------------------------------------------------------
        ws_usuarios = wb.create_sheet(title="Usuarios")
        ws_usuarios.append(["ID", "Nombre Completo", "Correo Electrónico", "Documento", "Rol", "Estado"])
        
        for u in datos.get('usuarios', []):
            ws_usuarios.append([
                u.get('id', 'N/A'),
                u.get('nombre', 'N/A'),
                u.get('correo', 'N/A'),
                u.get('documento', 'N/A'),
                u.get('rol', 'N/A'),
                u.get('estado', 'N/A')
            ])

        # ---------------------------------------------------------
        # 4. PESTAÑA: VEHÍCULOS
        # ---------------------------------------------------------
        ws_vehiculos = wb.create_sheet(title="Vehículos")
        ws_vehiculos.append(["ID", "Placa", "Tipo", "Propietario", "Documento Propietario", "Estado"])
        
        for v in datos.get('vehiculos', []):
            ws_vehiculos.append([
                v.get('id', 'N/A'),
                v.get('placa', 'N/A'),
                v.get('tipo', 'N/A'),
                v.get('propietario', 'N/A'),
                v.get('documento_propietario', 'N/A'),
                v.get('estado', 'N/A')
            ])

        # ---------------------------------------------------------
        # 5. PESTAÑA: RESERVAS
        # ---------------------------------------------------------
        ws_reservas = wb.create_sheet(title="Reservas")
        ws_reservas.append(["ID Reserva", "Usuario", "Placa", "Celda", "Fecha Reserva", "Estado"])
        
        for r in datos.get('reservas', []):
            ws_reservas.append([
                r.get('id', 'N/A'),
                r.get('usuario', 'N/A'),
                r.get('placa', 'N/A'),
                r.get('celda', 'N/A'),
                str(r.get('fecha_reserva', 'N/A')),
                r.get('estado', 'N/A')
            ])

        # ---------------------------------------------------------
        # FORMATO GENERAL Y AUTO-AJUSTE DE COLUMNAS
        # ---------------------------------------------------------
        for sheet in wb.worksheets:
            # Si no es la primera fila de título en Resumen, aplicar encabezado
            start_row = 4 if sheet.title == "Resumen General" else 1
            
            # Estilar encabezados de tabla
            for cell in sheet[start_row]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Aplicar bordes a las celdas con contenido
            for row in sheet.iter_rows(min_row=start_row, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
                for cell in row:
                    cell.border = border_thin
                    if cell.row > start_row:
                        cell.alignment = Alignment(vertical="center")

            # Ajustar ancho automático de cada columna
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Guardar en buffer RAM
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer


    @staticmethod
    def generar_pdf_completo(datos):
        """
        Genera un informe PDF corporativo de múltiples páginas con tablas dinámicas,
        resumen estadístico y el historial consolidado de ParkLink.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        
        # Estilos personalizados
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            textColor=colors.HexColor('#1A365D'),
            spaceAfter=4
        )
        
        subtitle_style = ParagraphStyle(
            'SubTitleStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#4A5568'),
            spaceAfter=15
        )

        section_heading = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12,
            textColor=colors.HexColor('#1A365D'),
            spaceBefore=12,
            spaceAfter=8
        )

        cell_style = ParagraphStyle(
            'TableCell',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            leading=10
        )

        cell_header_style = ParagraphStyle(
            'TableHeaderCell',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            leading=10,
            textColor=colors.white
        )

        story = []

        # 1. Encabezado
        story.append(Paragraph("PARKLINK - Informe Histórico del Sistema", title_style))
        story.append(Paragraph(f"Consolidado Global de Auditoría | Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E0'), spaceAfter=15))

        # 2. Tabla Resumen Metricas
        story.append(Paragraph("1. Resumen Estadístico Global", section_heading))
        
        data_resumen = [
            [Paragraph("Métrica", cell_header_style), Paragraph("Total Registros", cell_header_style)],
            [Paragraph("Usuarios Registrados", cell_style), Paragraph(str(len(datos.get('usuarios', []))), cell_style)],
            [Paragraph("Vehículos Registrados", cell_style), Paragraph(str(len(datos.get('vehiculos', []))), cell_style)],
            [Paragraph("Total Movimientos (Accesos)", cell_style), Paragraph(str(len(datos.get('accesos', []))), cell_style)],
            [Paragraph("Total Reservas Realizadas", cell_style), Paragraph(str(len(datos.get('reservas', []))), cell_style)],
        ]
        
        t_resumen = Table(data_resumen, colWidths=[3.0 * inch, 2.5 * inch])
        t_resumen.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A365D')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E0')),
        ]))
        story.append(t_resumen)
        story.append(Spacer(1, 15))

        # 3. Tabla Historial de Accesos (Movimientos)
        story.append(Paragraph("2. Historial Reciente de Accesos (Entradas y Salidas)", section_heading))

        headers_accesos = ["ID", "Usuario", "Placa", "Tipo", "Entrada", "Salida", "Estado"]
        data_accesos = [[Paragraph(h, cell_header_style) for h in headers_accesos]]

        accesos_list = datos.get('accesos', [])
        # Mostramos hasta 50 registros en el PDF para evitar un documento masivo
        for acc in accesos_list[:50]:
            data_accesos.append([
                Paragraph(str(acc.get('id', '')), cell_style),
                Paragraph(str(acc.get('usuario', 'N/A')), cell_style),
                Paragraph(str(acc.get('placa', '')), cell_style),
                Paragraph(str(acc.get('tipo_vehiculo', '')), cell_style),
                Paragraph(str(acc.get('fecha_entrada', '')), cell_style),
                Paragraph(str(acc.get('fecha_salida', 'En sitio')) if acc.get('fecha_salida') else 'En sitio', cell_style),
                Paragraph(str(acc.get('estado', '')), cell_style)
            ])

        if len(data_accesos) == 1:
            data_accesos.append([Paragraph("No hay registros de acceso disponibles.", cell_style)] + [Paragraph("", cell_style)] * 6)

        t_accesos = Table(data_accesos, colWidths=[0.4*inch, 1.8*inch, 0.8*inch, 0.8*inch, 1.3*inch, 1.3*inch, 0.8*inch])
        t_accesos.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A365D')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7FAFC')])
        ]))
        
        story.append(t_accesos)

        # Construir el PDF
        doc.build(story)
        buffer.seek(0)
        return buffer